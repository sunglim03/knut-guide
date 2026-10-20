"""맛집 엑셀 -> 앱 데이터(data/restaurants.js).

사용법:  python tools/build_data.py "엑셀 경로"
- 엑셀의 값은 그대로 옮긴다('확인 안됨'도 그대로. 화면에서 '방문 전 확인 필요'로 보여 준다).
- 지도 핀 좌표는 엑셀의 지도 링크에 접속해 가져오고 data/coords.json에 저장해 둔다(다음부터는 다시 접속하지 않음).
- 링크로 좌표를 못 얻은 가게는, 같은 건물 주소(도로명+번지)의 다른 가게 좌표를 쓴다. 그것도 없으면 지도에 찍지 않는다.
"""
import json, re, sys, time, urllib.error, urllib.parse, urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / "data" / "coords.json"
OUT = ROOT / "data" / "restaurants.js"
NA = "확인 안됨"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36",
      "Accept-Language": "ko-KR,ko;q=0.9"}
KAKAO_H = dict(UA, pf="web", appVersion="6.6.0", Referer="https://place.map.kakao.com/", Origin="https://place.map.kakao.com")


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *a, **k):
        return None


opener = urllib.request.build_opener(NoRedirect)


def final_url(url):
    """짧은 링크(naver.me 등)를 따라가 최종 주소를 얻는다."""
    for _ in range(6):
        if re.search(r"place/\d+|place\.map\.kakao\.com/\d+|[?&]lat=", url):
            return url
        try:
            opener.open(urllib.request.Request(url, headers=UA), timeout=15)
            return url
        except urllib.error.HTTPError as e:
            loc = e.headers.get("Location")
            if e.code in (301, 302, 303, 307, 308) and loc:
                url = urllib.parse.urljoin(url, loc)
                continue
            return url
    return url


def get(url, headers=UA):
    with urllib.request.urlopen(urllib.request.Request(url, headers=headers), timeout=20) as r:
        return r.read().decode("utf-8", "replace")


def coords_from_link(link):
    url = final_url(link)
    m = re.search(r"place\.map\.kakao\.com/(\d+)", url)
    if m:
        p = json.loads(get("https://place-api.map.kakao.com/places/panel3/" + m.group(1), KAKAO_H))["summary"]["point"]
        return p["lat"], p["lon"], "카카오맵 링크"
    m = re.search(r"place/(\d+)", url)
    if m:
        for attempt in range(4):
            try:
                html = get("https://m.place.naver.com/place/%s/home" % m.group(1))
                break
            except urllib.error.HTTPError as e:
                if e.code != 429:
                    raise
                time.sleep(10 + attempt * 10)
        else:
            return None
        c = re.search(r'"coordinate":\{[^}]*"x":"([\d.]+)","y":"([\d.]+)"', html) or re.search(r'"x":"([\d.]+)","y":"([\d.]+)"', html)
        return (float(c.group(2)), float(c.group(1)), "네이버 지도 링크") if c else None
    m = re.search(r"[?&]lng=([\d.]+).*?[?&]lat=([\d.]+)", url)
    if m:
        return float(m.group(2)), float(m.group(1)), "네이버 지도 링크(좌표)"
    return None  # 검색 링크 등


# 음식 종류(엑셀에 없는 칸): 엑셀 '메뉴'를 보고 분류. 메뉴가 '확인 안됨'이면 가게 이름에 음식이 드러날 때만 이름으로 분류.
# 근거가 없으면 '기타'. 카페·빵·술집·치킨·양식·학식도 '기타'. 새 가게가 생기면 여기에 추가하세요.
CATEGORY = {
    "한식": {"메뉴": ["국대부대찌개 교통대점", "부대담", "두꺼비네", "막국수파는집-한국교통대점", "배가네국수", "새마을 삼겹살",
                   "생각나는감자탕", "윤박사한식삼겹살", "일미닭갈비 교통대점", "제갈덮밥&곱창마니 한국교통대점", "한솥도시락 충주교통대점",
                   "24시전주명가콩나물국밥", "뚝딱김치찌개삼겹살", "만선식당", "메밀마당 충주 중앙탑본점", "몽돌집", "성마루 국밥",
                   "쌍둥이칼국수", "옛날순대", "오복쌈밥", "이정동묵밥", "일번지순대국", "제주돼지촌", "중원육면 충주점",
                   "촌집돌판삼겹살", "큰집아저씨 충주 본점"],
           "가게 이름": ["보스찜닭", "원조중원탑막국수", "옹심촌 감자옹심이 전문점", "교통대대패삼겹살 교통대본점", "큰집막창 봉방점"]},
    "중식": {"메뉴": ["리밍 교통대점", "관현각 본점", "권짬뽕", "탄금반점", "마라야 교통대점", "신신마라탕 교통대점"],
           "가게 이름": ["사천마라탕교통대점", "짬뽕생각", "차이나김쉐프"]},
    "일식": {"메뉴": ["도란도락도시락 본점", "타베", "막부", "봉방돈까스", "빨간사과", "쇼쿠스시 충주호암점", "예반72",
                   "우동키노야 충주호암점", "후라토식당 충주점", "핵밥 교통대점"],
           "가게 이름": ["카츠집", "돈까스카페"]},
    "분식": {"메뉴": ["대학로 친구분식", "미미분식", "신전떡볶이 충주교통대점", "오뚜기분식", "이삭토스트 충주한국교통대점",
                   "봉구스밥버거 교통대점"],
           "가게 이름": ["성대분식", "엄정분식"]},
}


def category_of(name):
    for cat, by in CATEGORY.items():
        for basis, names in by.items():
            if name in names:
                return cat, basis
    return "기타", None


def building_key(addr):
    m = re.search(r"([가-힣0-9]+(?:로|길))\s*(\d+(?:-\d+)?)", addr or "")
    return m.group(1) + " " + m.group(2) if m else None


def main(xlsx):
    ws = openpyxl.load_workbook(xlsx, data_only=True).worksheets[0]
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head) if h}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[col["이름"]]:
            continue  # 합계 줄 등
        g = lambda h: (str(r[col[h]]).strip() if r[col[h]] is not None else NA)
        rows.append({"name": g("이름"), "address": g("위치"), "link": g("지도 링크"), "category": g("분류"),
                     "checked": g("확인한 날"), "daehakro": g("충북 충주시 대소원면 대학로 위치 여부"),
                     "phone": g("전화번호"), "menu": g("메뉴"), "count": int(r[col["중복된 갯수"]] or 1)})

    cache = json.loads(CACHE.read_text(encoding="utf-8")) if CACHE.exists() else {}
    for p in rows:
        key = p["name"] + "|" + p["link"]
        if key in cache or not p["link"].startswith("http"):
            continue
        try:
            c = coords_from_link(p["link"])
        except Exception as e:  # noqa
            print("  좌표 실패:", p["name"], e)
            c = None
        cache[key] = {"lat": c[0], "lng": c[1], "source": c[2]} if c else None
        print("  %-24s %s" % (p["name"], cache[key]))
        time.sleep(1.2)
    CACHE.parent.mkdir(parents=True, exist_ok=True)
    CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")

    by_building = {}
    for p in rows:
        c = cache.get(p["name"] + "|" + p["link"])
        if c and building_key(p["address"]):
            by_building.setdefault(building_key(p["address"]), c)
    for p in rows:
        c = cache.get(p["name"] + "|" + p["link"])
        if not c and building_key(p["address"]) in by_building:
            b = by_building[building_key(p["address"])]
            c = {"lat": b["lat"], "lng": b["lng"], "source": "같은 건물(%s)의 다른 가게 위치" % building_key(p["address"])}
        p["lat"], p["lng"], p["coordSource"] = (c["lat"], c["lng"], c["source"]) if c else (None, None, None)

    names = {p["name"] for p in rows}
    unknown = [n for by in CATEGORY.values() for ns in by.values() for n in ns if n not in names]
    if unknown:
        print("  분류표에 있지만 엑셀에 없는 이름:", unknown)
    for p in rows:
        p["food"], p["foodBasis"] = category_of(p["name"])

    OUT.write_text("// tools/build_data.py로 생성. 직접 고치지 말고 엑셀을 고친 뒤 다시 실행하세요.\n"
                   "window.RESTAURANTS = " + json.dumps(rows, ensure_ascii=False, indent=1) + ";\n", encoding="utf-8")
    on_map = sum(1 for p in rows if p["lat"])
    print("가게 %d곳 / 지도 표시 %d곳 / 좌표 없음 %d곳 / 추천 수 합 %d" % (len(rows), on_map, len(rows) - on_map, sum(p["count"] for p in rows)))
    print("음식 종류:", {c: sum(1 for p in rows if p["food"] == c) for c in ["한식", "중식", "일식", "분식", "기타"]})
    for p in rows:
        if not p["lat"]:
            print("  지도에 없음:", p["name"], "|", p["address"])
        elif p["coordSource"].startswith("같은 건물"):
            print("  같은 건물 좌표:", p["name"], "|", p["coordSource"])


if __name__ == "__main__":
    main(sys.argv[1])
